import base64
import csv
import io
import zipfile
from datetime import date, datetime

from .fleet_domain import fail

HEADERS = {
    "Registration No": "registrationNumber",
    "Brand": "brand",
    "Model": "model",
    "Year": "year",
    "Color": "color",
    "VIN": "vin",
    "Engine No": "engineNumber",
    "Category": "category",
    "Daily Rate": "dailyRate",
    "Mileage": "mileage",
    "Owner": "ownerName",
    "Purchase Value": "purchaseValue",
    "Purchase Date": "purchaseDate",
}


def parse_import(payload, state):
    try:
        raw = base64.b64decode(payload.get("content", ""), validate=True)
        if len(raw) > 700000:
            fail("Import file must be smaller than 700 KB.")
        if str(payload.get("name", "")).lower().endswith(".csv"):
            rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
        elif str(payload.get("name", "")).lower().endswith(".xlsx"):
            from openpyxl import load_workbook

            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if (
                    sum(x.file_size for x in archive.infolist()) > 10_000_000
                    or len(archive.infolist()) > 100
                ):
                    fail("Spreadsheet expands beyond the supported size.")
            workbook = load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            rows = []
            try:
                for row in workbook.active.iter_rows(values_only=True):
                    rows.append(list(row))
                    if len(rows) > 501:
                        fail("Import at most 500 rows.")
            finally:
                workbook.close()
        else:
            fail("Choose a CSV or XLSX file.")
    except (ValueError, UnicodeError, zipfile.BadZipFile, KeyError, OSError):
        fail("The spreadsheet could not be read. Use the supplied template.")
    if not rows or len(rows) > 501:
        fail("Include a header and up to 500 vehicle rows.")
    headers = [HEADERS.get(str(x).strip(), str(x).strip()) for x in rows[0]]
    if not {"registrationNumber", "brand", "model", "year", "dailyRate", "ownerName"} <= set(
        headers
    ):
        fail("Required columns are missing. Download and use the template.")
    existing = {
        "".join(v["registrationNumber"].upper().split()) for v in state["vehicles"].values()
    }
    result = []
    for values in rows[1:]:
        if not any(v is not None and str(v).strip() for v in values):
            continue
        row = {
            k: v if v is not None else "" for k, v in zip(headers, values) if k in HEADERS.values()
        }
        errors = []
        for key in ("registrationNumber", "brand", "model", "ownerName"):
            row[key] = str(row.get(key, "")).strip()
            if not row[key]:
                errors.append(f"Missing {key}")
        for key in ("year", "dailyRate", "mileage", "purchaseValue"):
            try:
                value = float(row.get(key) or 0)
                if value < 0 or value > 1_000_000_000:
                    raise ValueError()
                row[key] = int(value) if key in {"year", "mileage"} else value
                if key == "year" and not 1950 <= value <= 2100:
                    raise ValueError()
            except (ValueError, TypeError, OverflowError):
                errors.append(f"Invalid {key}")
                row[key] = 0
        if isinstance(row.get("purchaseDate"), (date, datetime)):
            row["purchaseDate"] = row["purchaseDate"].isoformat()[:10]
        if not row.get("purchaseDate"):
            row.pop("purchaseDate", None)
        for key in ("color", "vin", "engineNumber", "category"):
            row[key] = str(row.get(key, "")).strip()
        norm = "".join(row["registrationNumber"].upper().split())
        if norm in existing:
            errors.append("Duplicate registration")
        existing.add(norm)
        if not any(
            o["name"].casefold() == row["ownerName"].casefold() for o in state["owners"].values()
        ):
            errors.append("Create this owner before importing")
        if not any(c["slug"] == row["category"] for c in state["categories"].values()):
            errors.append("Unknown vehicle category")
        row.update(isValid=not errors, errors=errors)
        result.append(row)
    return result
